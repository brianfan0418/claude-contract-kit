"""Native Word fixture: revisions, anchored comments, resolve metadata and round trip."""
import io
import json
from pathlib import Path
import subprocess
import tempfile
import unittest
import zipfile
from docx import Document
from lxml import etree as E
if __package__:
    from .tools import word_review as word
else:
    import importlib.util
    spec=importlib.util.spec_from_file_location("native_word_review",Path(__file__).parent/"tools"/"word_review.py")
    word=importlib.util.module_from_spec(spec)
    spec.loader.exec_module(word)


def fixture():
    d=Document();d.add_heading('虛構範例合約',0);d.add_heading('第 1 條 付款',1);p=d.add_paragraph('付款應於三十日內完成。');d.add_comment(p.runs,text='建議核對付款期限。',author='範例審閱人甲',initials='甲');d.add_heading('第 2 條 通知',1);d.add_paragraph('通知應以書面送達。');buf=io.BytesIO();d.save(buf)
    raw=word.suggest(buf.getvalue(),[{'type':'replace','clause_id':'2','quote':'書面','replacement':'書面或電子郵件'}],author='範例審閱人乙')
    parts=word.package(raw);c=word.xml(parts['word/comments.xml']);c[0].find('w:p',word.NS).set('{'+word.W14+'}paraId','AABBCCDD');parts['word/comments.xml']=word.dump_xml(c)
    rels=word.xml(parts['word/_rels/document.xml.rels']);n=E.SubElement(rels,'{'+word.R+'}Relationship');n.set('Id','rIdCommentsExtended');n.set('Type','http://schemas.microsoft.com/office/2011/relationships/commentsExtended');n.set('Target','commentsExtended.xml');parts['word/_rels/document.xml.rels']=word.dump_xml(rels)
    ct=word.xml(parts['[Content_Types].xml']);n=E.SubElement(ct,'{'+word.CT+'}Override');n.set('PartName','/word/commentsExtended.xml');n.set('ContentType','application/vnd.ms-word.commentsExtended+xml');parts['[Content_Types].xml']=word.dump_xml(ct)
    parts['word/commentsExtended.xml']=f'<w15:commentsEx xmlns:w15="{word.W15}"><w15:commentEx w15:paraId="AABBCCDD" w15:done="1"/></w15:commentsEx>'.encode()
    buf=io.BytesIO()
    with zipfile.ZipFile(buf,'w',zipfile.ZIP_DEFLATED) as z:
        for name,value in parts.items():z.writestr(name,value)
    return buf.getvalue()


class NativeWordTests(unittest.TestCase):
    def test_native_revisions_and_comment_anchor(self):
        r=word.inspect(fixture());self.assertEqual([c['id'] for c in r['clauses']],['1','2']);self.assertEqual(r['clauses'][1]['text'],'通知應以書面或電子郵件送達。')
        c=next(i for i in r['word_items'] if i['type']=='comment');self.assertEqual(c['author'],'範例審閱人甲');self.assertEqual(c['quote'],'付款應於三十日內完成。');self.assertEqual(c['clause_id'],'1');self.assertTrue(c['resolved'])
        self.assertEqual([(i['type'],i['quote']) for i in r['word_items'] if i['type']!='comment'],[('delete','書面'),('insert','書面或電子郵件')]);self.assertEqual(r['unresolved'],2)
    def test_ai_comment_is_native_and_unrelated_parts_preserved(self):
        raw=fixture();out=word.suggest(raw,[{'type':'comment','clause_id':'1','quote':'三十日','content':'請確認現金流安排。'}],author='AI（Codex／Claude）');before=word.package(raw);after=word.package(out)
        for name,value in before.items():
            if name not in ('word/document.xml','word/comments.xml','word/_rels/document.xml.rels','[Content_Types].xml'):self.assertEqual(after[name],value,name)
        comments=[i for i in word.inspect(out)['word_items'] if i['type']=='comment'];self.assertEqual(len(comments),2);self.assertEqual(comments[-1]['quote'],'三十日');self.assertFalse(comments[-1]['resolved']);self.assertEqual(comments[-1]['author'],'AI（Codex／Claude）')
        self.assertEqual(len(Document(io.BytesIO(out)).comments),2)
    def test_ambiguous_or_revision_anchor_refuses(self):
        with self.assertRaises(ValueError):word.suggest(fixture(),[{'type':'replace','clause_id':'2','quote':'書面或電子郵件','replacement':'電話'}])
        with self.assertRaises(ValueError):word.suggest(fixture(),[{'type':'comment','clause_id':'1','quote':'不存在','content':'建議'}])
    def test_invalid_file_or_empty_proposal_refuses(self):
        with self.assertRaises(ValueError):word.inspect(b'not docx')
        with self.assertRaises(ValueError):word.suggest(fixture(),[])
        with self.assertRaises(ValueError):word.package(b'x'*(word.MAX_BYTES+1))
    def test_record_keeps_word_and_schema(self):
        import base64
        raw=fixture();r=word.record(raw,'CASE-DEMO','/private/review.docx','範例甲');self.assertEqual(r['filename'],'review.docx');self.assertEqual(base64.b64decode(r['document_base64']),raw);self.assertEqual(r['format'],'docx');self.assertEqual(r['case_id'],'CASE-DEMO')
    def test_cli_output_is_exclusive_and_roundtrips(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);source=root/'original.docx';original=fixture();source.write_bytes(original);proposal=root/'ai.json';proposal.write_text(json.dumps([{'type':'comment','clause_id':'1','quote':'付款','content':'範例建議'}]));out=root/'new.docx';command=['python3',str(Path(word.__file__)),'suggest',str(source),'--suggestions',str(proposal),'--out',str(out)]
            first=subprocess.run(command,capture_output=True);self.assertEqual(first.returncode,0,first.stderr);self.assertEqual(source.read_bytes(),original);saved=out.read_bytes();second=subprocess.run(command,capture_output=True);self.assertNotEqual(second.returncode,0);self.assertEqual(out.read_bytes(),saved)
